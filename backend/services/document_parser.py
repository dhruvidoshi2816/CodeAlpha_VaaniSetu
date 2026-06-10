"""
document_parser.py — Extract plain text from TXT, PDF, and DOCX uploads.
"""

from __future__ import annotations

import io
import logging

from PyPDF2 import PdfReader
from PyPDF2.errors import PdfReadError
from docx import Document

logger = logging.getLogger(__name__)

MAX_PAGES = 500
_TEXT_ENCODINGS = ("utf-8", "utf-8-sig", "utf-16", "latin-1", "cp1252", "cp1251")


def _reset_stream(file_storage) -> None:
    stream = getattr(file_storage, "stream", file_storage)
    if hasattr(stream, "seek"):
        stream.seek(0)


def _read_bytes(file_storage) -> bytes:
    _reset_stream(file_storage)
    raw = file_storage.read()
    _reset_stream(file_storage)
    return raw


def extract_text_from_document(file_storage, filename: str) -> str:
    """Extract text from uploaded document. Raises ValueError on unsupported types."""
    if not filename or "." not in filename:
        raise ValueError("Filename must include an extension")

    ext = filename.rsplit(".", 1)[-1].lower()

    if ext == "txt":
        return _extract_txt(file_storage)
    if ext == "pdf":
        return _extract_pdf(file_storage)
    if ext == "docx":
        return _extract_docx(file_storage)

    raise ValueError(f"Unsupported file type: {ext}")


def _extract_txt(file_storage) -> str:
    raw = _read_bytes(file_storage)
    for encoding in _TEXT_ENCODINGS:
        try:
            return raw.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace").strip()


def _extract_pdf(file_storage) -> str:
    raw = _read_bytes(file_storage)
    try:
        reader = PdfReader(io.BytesIO(raw))
    except PdfReadError as exc:
        raise ValueError(f"Could not read PDF: {exc}") from exc

    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ValueError(
                "PDF is password-protected. Please upload an unlocked PDF."
            ) from exc

    pages: list[str] = []
    for i, page in enumerate(reader.pages):
        if i >= MAX_PAGES:
            logger.warning("PDF truncated at %d pages", MAX_PAGES)
            break
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            logger.warning("PDF page %d extract failed: %s", i + 1, exc)
            text = ""
        text = text.strip()
        if text:
            pages.append(text)

    if not pages:
        raise ValueError(
            "No text found in PDF. This may be a scanned image PDF — "
            "try uploading a text-based PDF or use OCR on an image export."
        )

    return "\n\n".join(pages)


def _extract_docx(file_storage) -> str:
    _reset_stream(file_storage)
    try:
        doc = Document(file_storage)
    except Exception as exc:
        raise ValueError(f"Could not read DOCX file: {exc}") from exc
    finally:
        _reset_stream(file_storage)

    parts: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    result = "\n\n".join(parts).strip()
    if not result:
        raise ValueError("No text found in DOCX document")
    return result
